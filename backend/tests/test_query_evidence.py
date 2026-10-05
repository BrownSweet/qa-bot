"""A real application DB retains the source and query facts across a chat stream."""
import asyncio
import csv
import io
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import analysis_tasks, database, models, schemas, semantic_context
from app.config import settings
from app.database import Base
from app.routers import chat, export, sessions
from app.source_identity import source_identity_json


@pytest.fixture
def owned_db(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool,
                           connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    monkeypatch.setattr(database, "SessionLocal", factory)
    db = factory()
    user = models.User(id="owner", username="owner", phone="19900000001", password_hash="hash")
    outsider = models.User(id="other", username="other", phone="19900000002", password_hash="hash")
    db.add_all([user, outsider])
    db.flush()
    source_a = models.DBConfig(id="source-a", user_id=user.id, name="账本 A", type="sqlite",
                               file_path="/synthetic/a.sqlite")
    source_b = models.DBConfig(id="source-b", user_id=user.id, name="账本 B", type="sqlite",
                               file_path="/synthetic/b.sqlite")
    db.add_all([
        source_a, source_b,
        models.DBConfig(id="foreign-source", user_id=outsider.id, name="他人", type="sqlite",
                        file_path="/synthetic/private.sqlite"),
        models.ChatSession(id="session", user_id=user.id, name="分析"),
        models.Message(id="old-a", session_id="session", user_id=user.id,
                       role="user", content="只属于 A 的上下文", status="completed",
                       db_config_id="source-a"),
        semantic_context.SourceSemantics(id="meaning-b", user_id=user.id,
                                         db_config_id="source-b", context="收入以已支付订单计算",
                                         source_identity_json=source_identity_json(source_b)),
        semantic_context.SourceSemantics(id="meaning-a", user_id=user.id,
                                         db_config_id="source-a", context="A 专属口径",
                                         source_identity_json=source_identity_json(source_a)),
    ])
    db.commit()
    yield db, user, outsider
    db.close()
    engine.dispose()


def test_source_is_bound_and_evidence_survives_stream(owned_db, monkeypatch):
    db, user, _ = owned_db
    captured = {}
    synthetic_engine = SimpleNamespace(dialect=SimpleNamespace(name="sqlite"), dispose=lambda: None)
    monkeypatch.setattr(chat, "validate_data_file", lambda path, kind: path)
    monkeypatch.setattr(chat.engine_utils, "build_engine", lambda config: synthetic_engine)
    monkeypatch.setattr(chat.engine_utils, "get_schema", lambda engine: "orders(amount INTEGER)")
    monkeypatch.setattr(chat.engine_utils, "run_query_with_metadata", lambda engine, sql: (
        ["amount"], [["=1+1"], [20]],
        {"row_limit": 200, "returned_rows": 2, "has_more": False},
    ))
    monkeypatch.setattr(chat.ai, "get_ai_config", lambda db: {"api_key": "synthetic"})

    async def sql(_cfg, _schema, _question, **kwargs):
        captured.update(kwargs)
        return "SELECT amount FROM orders"

    async def answer(*_args, **_kwargs):
        yield "本次有两行"

    monkeypatch.setattr(chat.ai, "generate_sql", sql)
    monkeypatch.setattr(chat.ai, "analyze_stream", answer)
    response = chat.send_message(schemas.SendMessageRequest(
        session_id="session", db_config_id="source-b", question="总额？",
    ), db=db, user=user)

    async def consume():
        return [event async for event in response.body_iterator]

    events = asyncio.run(consume())
    assert captured["history"] == []
    assert captured["business_context"] == "收入以已支付订单计算"
    result_event = next(event for event in events if event.startswith("event: result"))
    assert '"rows": [["=1+1"], [20]]' in result_event
    assert '"status": "completed"' in events[-1]

    db.expire_all()
    assert db.get(models.ChatSession, "session").db_config_id == "source-b"
    messages = chat.get_messages("session", db=db, user=user)["messages"]
    answer_message = next(message for message in messages if message["role"] == "assistant")
    question_message = next(message for message in messages if message["content"] == "总额？")
    assert answer_message["question_message_id"] == question_message["id"]
    assert answer_message["source"] == {
        "id": "source-b", "name": "账本 B", "type": "sqlite",
        "host": None, "port": None, "database": None, "username": None,
        "file_path": "/synthetic/b.sqlite",
    }
    assert answer_message["evidence"]["sql"] == "SELECT amount FROM orders"
    assert answer_message["evidence"]["generated_sql"] == "SELECT amount FROM orders"
    assert answer_message["evidence"]["rows"] == [["=1+1"], [20]]
    assert answer_message["evidence"]["coverage"]["evidence_truncated"] is False
    assert answer_message["evidence"]["executed_at"].endswith("Z")
    snapshot = answer_message["generation_snapshot"]
    assert snapshot["question"] == "总额？"
    assert snapshot["schema"] == "orders(amount INTEGER)"
    assert snapshot["business_context"] == "收入以已支付订单计算"
    assert snapshot["history"] == [] and snapshot["task"] is None
    assert snapshot["api_url_sha256"] and "api_url" not in snapshot
    complete = json.loads(events[-1].split("data: ", 1)[1])
    assert complete["generation_snapshot"] == snapshot


def test_failed_sql_generation_retains_verified_task_inputs_without_web_endpoint(owned_db, monkeypatch):
    db, user, _ = owned_db
    source = db.get(models.DBConfig, "source-b")
    task = analysis_tasks.AnalysisTask(
        id="snapshot-task", user_id=user.id, db_config_id=source.id,
        name="月销售", question_template="查看{{month}}销售额",
        source_identity_json=source_identity_json(source),
    )
    db.add(task)
    db.commit()
    monkeypatch.setattr(settings, "DESKTOP_MODE", False)
    monkeypatch.setattr(chat, "validate_data_file", lambda path, _kind: path)
    synthetic_engine = SimpleNamespace(dialect=SimpleNamespace(name="sqlite"), dispose=lambda: None)
    monkeypatch.setattr(chat.engine_utils, "build_engine", lambda _cfg: synthetic_engine)
    monkeypatch.setattr(chat.engine_utils, "get_schema", lambda _engine: "sales(month TEXT, amount INTEGER)")
    monkeypatch.setattr(chat.ai, "get_ai_config", lambda _db: {
        "api_key": "synthetic-secret", "api_url": "https://internal.example/v1",
        "model": "internal-model",
    })

    async def failed_sql(*_args, **_kwargs):
        # The provider has not been contacted until this database commit exists.
        with database.SessionLocal() as fresh:
            saved = fresh.query(models.Message).filter(models.Message.role == "assistant").one()
            assert json.loads(saved.generation_snapshot_json)["task"]["parameters"] == {"month": "九月"}
        raise RuntimeError("synthetic provider failure")

    monkeypatch.setattr(chat.ai, "generate_sql", failed_sql)
    response = chat.send_message(schemas.SendMessageRequest(
        session_id="session", db_config_id=source.id, question="查看九月销售额",
        analysis_task_id=task.id, analysis_task_parameters={"month": "九月"},
    ), db=db, user=user)

    async def consume():
        return [event async for event in response.body_iterator]

    events = asyncio.run(consume())
    answer = next(message for message in chat.get_messages("session", db=db, user=user)["messages"]
                  if message["role"] == "assistant")
    snapshot = answer["generation_snapshot"]
    assert answer["status"] == "error" and answer["evidence"] is None
    assert snapshot["model"] == "internal-model"
    assert snapshot["schema"] == "sales(month TEXT, amount INTEGER)"
    assert snapshot["question"] == "查看九月销售额"
    assert snapshot["business_context"] == "收入以已支付订单计算"
    assert snapshot["history"] == []
    assert snapshot["task"] == {
        "id": task.id, "name": "月销售", "question_template": "查看{{month}}销售额",
        "parameters": {"month": "九月"},
    }
    serialized = json.dumps(snapshot, ensure_ascii=False)
    assert "internal.example" not in serialized and "synthetic-secret" not in serialized
    assert len(snapshot["api_url_sha256"]) == 64 and "api_url" not in snapshot
    assert json.loads(events[-1].split("data: ", 1)[1])["generation_snapshot"] == snapshot
    task.name = "已改名"
    task.question_template = "其他问题"
    db.commit()
    assert analysis_tasks.list_runs(task.id, db=db, user=user)["runs"][0]["generation_snapshot"] == snapshot
    csv_rows = list(csv.reader(io.StringIO(export.export_session(schemas.ExportRequest(
        session_id="session", format="csv",
    ), db=db, user=user).body.decode("utf-8-sig"))))
    assert json.loads(next(row for row in csv_rows if row[4] == answer["id"])[12]) == snapshot


def test_session_source_update_checks_owner(owned_db):
    db, user, _ = owned_db
    created = sessions.create_session(schemas.CreateSessionRequest(
        name="绑定", db_config_id="source-b",
    ), db=db, user=user)["session"]
    assert created["db_config_id"] == "source-b"
    listed = sessions.list_sessions(db=db, user=user)["sessions"]
    assert next(item for item in listed if item["id"] == created["id"])["db_config_id"] == "source-b"
    with pytest.raises(Exception) as invalid:
        sessions.update_session(created["id"], schemas.UpdateSessionRequest(
            db_config_id="foreign-source",
        ), db=db, user=user)
    assert invalid.value.status_code == 404
    assert sessions.update_session(created["id"], schemas.UpdateSessionRequest(
        db_config_id=None,
    ), db=db, user=user)["db_config_id"] is None


def test_task_runs_require_matching_saved_template_and_source(owned_db, monkeypatch):
    db, user, _ = owned_db
    task = analysis_tasks.AnalysisTask(
        id="task", user_id=user.id, db_config_id="source-b",
        name="月销售", question_template="查看{{month}}销售额",
        source_identity_json=source_identity_json(db.get(models.DBConfig, "source-b")),
    )
    db.add(task)
    db.commit()
    monkeypatch.setattr(chat, "validate_data_file", lambda path, kind: path)
    monkeypatch.setattr(chat.ai, "get_ai_config", lambda _db: {"api_key": "synthetic"})

    for body, status in [
        (schemas.SendMessageRequest(session_id="session", db_config_id="source-b",
                                    question="任意问题", analysis_task_id="task",
                                    analysis_task_parameters={"month": "九月"}), 400),
        (schemas.SendMessageRequest(session_id="session", db_config_id="source-b",
                                    question="查看九月销售额", analysis_task_id="task"), 400),
        (schemas.SendMessageRequest(session_id="session", db_config_id="source-a",
                                    question="查看九月销售额", analysis_task_id="task",
                                    analysis_task_parameters={"month": "九月"}), 404),
    ]:
        with pytest.raises(Exception) as rejected:
            chat.send_message(body, db=db, user=user)
        assert rejected.value.status_code == status
    assert db.query(models.Message).filter(models.Message.analysis_task_id == "task").count() == 0

    chat.send_message(schemas.SendMessageRequest(
        session_id="session", db_config_id="source-b", question="查看九月销售额",
        analysis_task_id="task", analysis_task_parameters={"month": "九月"},
    ), db=db, user=user)
    run = db.query(models.Message).filter(models.Message.role == "assistant",
                                           models.Message.analysis_task_id == "task").one()
    assert run.question_message_id
    assert db.get(models.Message, run.question_message_id).analysis_task_id == "task"

    db.get(models.DBConfig, "source-b").file_path = "/synthetic/replaced.sqlite"
    db.commit()
    assert semantic_context.get_context(db, "source-b", user.id) == ""
    with pytest.raises(Exception) as stale:
        chat.send_message(schemas.SendMessageRequest(
            session_id="session", db_config_id="source-b", question="查看九月销售额",
            analysis_task_id="task", analysis_task_parameters={"month": "九月"},
        ), db=db, user=user)
    assert stale.value.status_code == 409


def test_chat_and_result_exports_keep_formula_text_literal(owned_db):
    db, user, outsider = owned_db
    evidence = {"sql": "SELECT note FROM orders", "columns": ["note"],
                "rows": [["=1+1"], [" +SUM(A1:A2)"], ["ordinary"]],
                "coverage": {"returned_rows": 3, "evidence_rows": 3,
                             "has_more": False, "evidence_truncated": False,
                             "analysis_truncated": False},
                "executed_at": "2026-10-03T00:00:00Z"}
    db.add(models.Message(id="formula-answer", session_id="session", user_id=user.id,
                          role="assistant", content="=1+1", status="completed",
                          db_config_id="source-b", evidence_json=json.dumps(evidence),
                          source_snapshot_json=json.dumps({"name": "账本 B", "type": "sqlite",
                                                           "file_path": "/synthetic/b.sqlite"})))
    db.commit()

    chat_csv = export.export_session(schemas.ExportRequest(
        session_id="session", format="csv",
    ), db=db, user=user).body.decode("utf-8-sig")
    chat_rows = list(csv.reader(io.StringIO(chat_csv)))
    assert chat_rows[0][:4] == ["角色", "内容", "状态", "时间"]
    answer_row = next(row for row in chat_rows if row[0] == "助手")
    assert answer_row[1] == "'=1+1"
    assert json.loads(answer_row[5])["name"] == "账本 B"
    assert answer_row[6] == "SELECT note FROM orders"
    assert answer_row[8] == "2026-10-03T00:00:00Z"
    assert json.loads(answer_row[9])["evidence_rows"] == 3
    assert json.loads(answer_row[11]) == evidence["rows"]
    legacy_row = next(row for row in chat_rows if row[1] == "只属于 A 的上下文")
    assert legacy_row[6:] == [""] * 7

    chat_xlsx = export.export_session(schemas.ExportRequest(
        session_id="session", format="excel",
    ), db=db, user=user).body
    workbook = load_workbook(io.BytesIO(chat_xlsx))
    cell = workbook.active["B3"]
    assert cell.value == "=1+1" and cell.data_type == "s"
    assert workbook.active["G3"].value == "SELECT note FROM orders"
    assert workbook.active["L3"].data_type == "s"
    assert json.loads(workbook.active["L3"].value) == evidence["rows"]
    detail_rows = list(workbook["查询证据"].values)
    assert detail_rows[1][1] == "元数据"
    assert json.loads(detail_rows[1][5])["coverage"]["evidence_rows"] == 3
    assert json.loads(detail_rows[2][5]) == ["=1+1"]
    assert workbook["范围说明"]["B1"].value.startswith("查询证据仅包含")

    result_csv = export.export_result(schemas.ExportResultRequest(
        message_id="formula-answer", format="csv",
    ), db=db, user=user).body.decode("utf-8-sig")
    result_rows = list(csv.reader(io.StringIO(result_csv)))
    assert result_rows[0] == ["note", "__qa_record_type", "__qa_scope"]
    assert result_rows[1][1] == "metadata"
    scope = json.loads(result_rows[1][2])
    assert scope["source"]["name"] == "账本 B"
    assert scope["coverage"]["analysis_truncated"] is False
    assert scope["coverage"]["evidence_truncated"] is False
    assert result_rows[2:] == [
        ["'=1+1", "data", ""], ["' +SUM(A1:A2)", "data", ""], ["ordinary", "data", ""],
    ]
    result_xlsx = export.export_result(schemas.ExportResultRequest(
        message_id="formula-answer", format="excel",
    ), db=db, user=user).body
    workbook = load_workbook(io.BytesIO(result_xlsx))
    assert workbook["查询结果预览"]["A2"].value == "=1+1"
    assert workbook["查询结果预览"]["A2"].data_type == "s"
    assert workbook["范围说明"]["B3"].value == "SELECT note FROM orders"
    assert workbook["范围说明"]["B5"].value == "账本 B"
    assert workbook["范围说明"]["B11"].value == "/synthetic/b.sqlite"

    with pytest.raises(Exception) as forbidden:
        export.export_result(schemas.ExportResultRequest(
            message_id="formula-answer", format="csv",
        ), db=db, user=outsider)
    assert forbidden.value.status_code == 404


def test_web_history_and_export_hide_desktop_restored_provider_url(owned_db, monkeypatch):
    db, user, _ = owned_db
    monkeypatch.setattr(settings, "DESKTOP_MODE", False)
    saved = {"model": "m", "api_url": "https://private.internal.example/v1",
             "dialect": "sqlite", "schema": "sales(amount INTEGER)",
             "business_context": "", "history": [], "question": "总额", "task": None}
    db.add(models.Message(id="restored-answer", session_id="session", user_id=user.id,
                          role="assistant", content="旧回答", status="completed",
                          generation_snapshot_json=json.dumps(saved)))
    db.commit()
    visible = next(message for message in chat.get_messages("session", db=db, user=user)["messages"]
                   if message["id"] == "restored-answer")["generation_snapshot"]
    assert "api_url" not in visible and len(visible["api_url_sha256"]) == 64
    csv_text = export.export_session(schemas.ExportRequest(
        session_id="session", format="csv",
    ), db=db, user=user).body.decode("utf-8-sig")
    assert "private.internal.example" not in csv_text
    row = next(row for row in csv.reader(io.StringIO(csv_text)) if row[4] == "restored-answer")
    assert json.loads(row[12]) == visible


def test_desktop_provider_path_token_never_enters_new_snapshot_or_exports(owned_db, monkeypatch):
    db, user, _ = owned_db
    monkeypatch.setattr(settings, "DESKTOP_MODE", True)
    monkeypatch.setattr(chat, "validate_data_file", lambda path, _kind: path)
    synthetic_engine = SimpleNamespace(dialect=SimpleNamespace(name="sqlite"), dispose=lambda: None)
    monkeypatch.setattr(chat.engine_utils, "build_engine", lambda _cfg: synthetic_engine)
    monkeypatch.setattr(chat.engine_utils, "get_schema", lambda _engine: "sales(amount INTEGER)")
    monkeypatch.setattr(chat.ai, "get_ai_config", lambda _db: {
        "api_key": "provider-key", "api_url": "https://proxy.example.com/tenant-secret-token/v1",
        "model": "desktop-model",
    })

    async def failed_sql(*_args, **_kwargs):
        raise RuntimeError("synthetic provider failure")

    monkeypatch.setattr(chat.ai, "generate_sql", failed_sql)
    response = chat.send_message(schemas.SendMessageRequest(
        session_id="session", db_config_id="source-b", question="总额？",
    ), db=db, user=user)

    async def consume():
        return [event async for event in response.body_iterator]

    events = asyncio.run(consume())
    answer = next(message for message in chat.get_messages("session", db=db, user=user)["messages"]
                  if message["role"] == "assistant")
    with database.SessionLocal() as fresh:
        raw = fresh.get(models.Message, answer["id"]).generation_snapshot_json
    assert "tenant-secret-token" not in raw
    assert "provider-key" not in raw
    assert answer["generation_snapshot"]["api_url"] == "https://proxy.example.com"
    assert len(answer["generation_snapshot"]["api_url_sha256"]) == 64
    assert "tenant-secret-token" not in events[-1]

    # Old Desktop records cannot be rewritten in an already distributed backup,
    # but every current read and export must suppress their path tokens.
    db.add(models.Message(id="legacy-secret-path", session_id="session", user_id=user.id,
                          role="assistant", content="旧回答", status="completed",
                          generation_snapshot_json=json.dumps({
                              "model": "old", "api_url": "https://proxy.example.com/old-secret-path/v1",
                              "dialect": "sqlite", "schema": "sales(amount INTEGER)",
                              "business_context": "", "history": [], "question": "旧问题", "task": None,
                          })))
    db.commit()
    legacy = next(message for message in chat.get_messages("session", db=db, user=user)["messages"]
                  if message["id"] == "legacy-secret-path")
    assert legacy["generation_snapshot"]["api_url"] == "https://proxy.example.com"
    assert len(legacy["generation_snapshot"]["api_url_sha256"]) == 64

    csv_text = export.export_session(schemas.ExportRequest(
        session_id="session", format="csv",
    ), db=db, user=user).body.decode("utf-8-sig")
    assert "tenant-secret-token" not in csv_text and "old-secret-path/v1" not in csv_text
    records = list(csv.reader(io.StringIO(csv_text)))
    for record in records[1:]:
        if record[4] in {answer["id"], "legacy-secret-path"}:
            assert json.loads(record[12])["api_url"] == "https://proxy.example.com"
    workbook = load_workbook(io.BytesIO(export.export_session(schemas.ExportRequest(
        session_id="session", format="excel",
    ), db=db, user=user).body))
    exported_text = " ".join(str(value) for sheet in workbook for row in sheet.values for value in row)
    assert "tenant-secret-token" not in exported_text
    assert "old-secret-path/v1" not in exported_text


def test_evidence_records_exact_prefix_and_declares_byte_cap():
    evidence = chat._query_evidence("SELECT payload FROM records", ["payload"],
                                    [["a" * 10], ["b" * 10]],
                                    {"row_limit": 200, "returned_rows": 2, "has_more": False})
    assert evidence["rows"] == [["a" * 10], ["b" * 10]]
    oversized = chat._query_evidence("SELECT payload FROM records", ["payload"],
                                     [["x" * (chat.MAX_EVIDENCE_BYTES + 1)]],
                                     {"row_limit": 200, "returned_rows": 1, "has_more": False})
    assert oversized["rows"] == []
    assert oversized["coverage"]["evidence_rows"] == 0
    assert oversized["coverage"]["evidence_truncated"] is True
    rendered = chat._query_evidence("SELECT amount LIMIT 201", ["amount"], [[10]],
                                    {"row_limit": 200, "returned_rows": 1, "has_more": False},
                                    generated_sql="SELECT amount")
    assert rendered["sql"] == "SELECT amount LIMIT 201"
    assert rendered["generated_sql"] == "SELECT amount"
    nonfinite = chat._query_evidence("SELECT amount", ["amount"], [[float("nan")]],
                                     {"row_limit": 200, "returned_rows": 1, "has_more": False})
    assert nonfinite["rows"] == [["nan"]]
    assert json.loads(chat._sse("result", {"evidence": nonfinite}).split("data: ", 1)[1])


def test_empty_result_csv_still_carries_scope_and_complex_excel_cells(owned_db):
    db, user, _ = owned_db
    evidence = {"sql": "SELECT payload FROM records WHERE 0", "columns": ["payload"],
                "rows": [], "coverage": {"returned_rows": 0, "has_more": False,
                                         "analysis_truncated": False,
                                         "evidence_rows": 0, "evidence_truncated": False},
                "executed_at": "2026-10-03T00:00:00Z"}
    db.add(models.Message(id="empty-result", session_id="session", user_id=user.id,
                          role="assistant", content="没有记录", status="completed",
                          evidence_json=json.dumps(evidence)))
    db.commit()
    response = export.export_result(schemas.ExportResultRequest(
        message_id="empty-result", format="csv",
    ), db=db, user=user)
    records = list(csv.reader(io.StringIO(response.body.decode("utf-8-sig"))))
    assert records[0] == ["payload", "__qa_record_type", "__qa_scope"]
    assert len(records) == 2
    assert records[1][1] == "metadata"
    assert json.loads(records[1][2])["coverage"]["returned_rows"] == 0

    from openpyxl import Workbook
    workbook = Workbook()
    count = export._write_excel_rows(workbook.active, [[{"formula": "=1+1"}], ["x" * 40000]])
    assert count == 1
    assert workbook.active["A1"].data_type == "s"
    assert workbook.active["A2"].value.endswith("（Excel 单元格已截断）")


def test_session_export_chunks_wide_preview_and_rejects_oversized_session(owned_db, monkeypatch):
    db, user, _ = owned_db
    wide_row = ["x" * 8000 for _ in range(5)]
    evidence = {"sql": "SELECT * FROM wide", "columns": ["a", "b", "c", "d", "e"],
                "rows": [wide_row], "coverage": {"evidence_rows": 1, "has_more": False},
                "executed_at": "2026-10-03T00:00:00Z"}
    db.add(models.Message(id="wide-answer", session_id="session", user_id=user.id,
                          role="assistant", content="宽表", status="completed",
                          evidence_json=json.dumps(evidence),
                          generation_snapshot_json=json.dumps({"schema": "s" * 40000})))
    db.commit()
    workbook = load_workbook(io.BytesIO(export.export_session(schemas.ExportRequest(
        session_id="session", format="excel",
    ), db=db, user=user).body))
    main = list(workbook["会话记录"].values)
    wide_main = next(row for row in main if row[4] == "wide-answer")
    assert wide_main[11].startswith("见「查询证据」页")
    assert wide_main[12].startswith("见「查询证据」页")
    fragments = [row for row in workbook["查询证据"].values
                 if row[0] == "wide-answer" and row[1] == "预览行"]
    assert len(fragments) == 2
    assert json.loads("".join(row[5] for row in fragments)) == wide_row
    generation_fragments = [row for row in workbook["查询证据"].values
                            if row[0] == "wide-answer" and row[1] == "生成输入快照"]
    assert json.loads("".join(row[5] for row in generation_fragments))["schema"] == "s" * 40000

    monkeypatch.setattr(export, "MAX_SESSION_EXPORT_BYTES", 100)
    with pytest.raises(Exception) as too_large:
        export.export_session(schemas.ExportRequest(session_id="session", format="csv"),
                              db=db, user=user)
    assert too_large.value.status_code == 413

    monkeypatch.setattr(export, "MAX_SESSION_EXPORT_BYTES", 8 * 1024 * 1024)
    monkeypatch.setattr(export, "MAX_SESSION_PREVIEW_ROWS", 0)
    with pytest.raises(Exception) as too_many_rows:
        export.export_session(schemas.ExportRequest(session_id="session", format="excel"),
                              db=db, user=user)
    assert too_many_rows.value.status_code == 413
